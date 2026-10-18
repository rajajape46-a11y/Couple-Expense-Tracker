"""Part-2 browser test: summary pages, Excel export/import with duplicates, recurring expenses."""
import re, subprocess, sys, shutil, json
import openpyxl
import datetime as dt_mod
from playwright.sync_api import sync_playwright, expect

BASE = "http://localhost:3100"
SHOTS = "/home/claude/e2e/shots"
WORK = "/home/claude/e2e/files"
V2 = "/mnt/user-data/outputs/Monthly expense log v2.xlsx"
V1 = "/mnt/user-data/uploads/Personal Expense/Monthly expense log.xlsx"
RECALC = "/root/.claude/skills/synced/e64cbee3-f0a8-4f74-9982-1b296cdd2468_ea01e10e-e8b9-44eb-853e-4bd2918bd975/xlsx/scripts/recalc.py"
results = []


def check(name, fn):
    try:
        fn()
        results.append(("PASS", name))
    except Exception as e:  # noqa
        results.append(("FAIL", f"{name}: {str(e)[:400]}"))


def sql(q):
    return subprocess.run(["su", "postgres", "-c", f'psql -tA -d cet_e2e -c "{q}"'],
                          capture_output=True, text=True).stdout.strip()


def eq(a, b):
    assert a == b, f"{a!r} != {b!r}"


def live_count():
    return sql("select count(*) from transactions where deleted_at is null")


def import_file(page, path):
    page.goto(BASE + "/settings/data")
    expect(page.get_by_text("Choose an Excel file (.xlsx)")).to_be_visible()
    page.wait_for_load_state("networkidle")
    page.set_input_files("input[type=file]", path)
    expect(page.get_by_test_id("import-preview")).to_be_visible(timeout=20000)
    counts = {}
    for label in ["New", "Duplicates", "Problems"]:
        box = page.get_by_test_id("import-preview").locator("div.rounded-xl.bg-canvas").filter(has_text=label)
        counts[label] = int(box.locator("p").first.inner_text())
    return counts


shutil.rmtree(WORK, ignore_errors=True)
subprocess.run(["mkdir", "-p", WORK])

with sync_playwright() as p:
    browser = p.chromium.launch()
    ctx = browser.new_context(viewport={"width": 390, "height": 844}, device_scale_factor=2,
                              is_mobile=True, has_touch=True, timezone_id="Australia/Melbourne",
                              accept_downloads=True)
    page = ctx.new_page()
    page.on("dialog", lambda d: d.accept())
    page.goto(BASE + "/login")
    page.fill("input[name=email]", "raja@example.com")
    page.fill("input[name=password]", "test1234")
    page.click("button[type=submit]")
    expect(page).to_have_url(BASE + "/")

    # ---------------------------------------------------------------- home & nav
    check("settle-up card removed from home", lambda: expect(page.get_by_text("Settle up")).to_have_count(0))
    check("bottom nav has Summary", lambda: expect(page.get_by_role("link", name="Summary", exact=True)).to_be_visible())

    # ---------------------------------------------------------------- summary
    page.get_by_role("link", name="Summary", exact=True).click()
    expect(page).to_have_url(re.compile("/summary$"))
    grid = page.get_by_test_id("summary-grid")
    check("summary grid has Sep and Oct columns", lambda: (
        expect(grid.get_by_role("link", name="Sep")).to_be_visible(),
        expect(grid.get_by_role("link", name="Oct")).to_be_visible()))
    check("year total matches the spreadsheet ($10,891)", lambda: expect(page.get_by_text("$10,891").first).to_be_visible())
    check("grid row for Housing & Rent shows $7,519", lambda: expect(grid.locator("tr").filter(has_text="Housing & Rent")).to_contain_text("$7,519"))
    page.screenshot(path=f"{SHOTS}/p2-01-summary.png", full_page=True)

    grid.get_by_role("link", name="Oct").click()
    expect(page).to_have_url(re.compile("/summary/2026-10$"))
    check("month page title", lambda: expect(page.get_by_role("heading", name="October 2026")).to_be_visible())
    check("month total $10,285.20", lambda: expect(page.get_by_text("$10,285.20")).to_be_visible())
    check("vs previous month shown", lambda: expect(page.get_by_text("September 2026: $606")).to_be_visible())
    page.locator("summary").filter(has_text="Housing & Rent").click()
    check("category drill-down shows sub-categories", lambda: expect(page.get_by_text("Bond & Rental Fees", exact=True)).to_be_visible())
    check("biggest expense listed", lambda: expect(page.get_by_role("link", name=re.compile("Rent & Bond"))).to_be_visible())
    page.screenshot(path=f"{SHOTS}/p2-02-month.png", full_page=True)

    # ---------------------------------------------------------------- export
    page.goto(BASE + "/settings/data")
    with page.expect_download() as dl:
        page.get_by_role("link", name="Export to Excel").click()
    exported = f"{WORK}/exported.xlsx"
    dl.value.save_as(exported)

    def verify_export():
        name = dl.value.suggested_filename
        assert re.match(r"Monthly expense log \d{4}-\d{2}-\d{2}\.xlsx", name), name
        wb = openpyxl.load_workbook(exported)
        assert wb.sheetnames == ["Dashboard", "Expense Log", "Monthly Summary", "Lists"], wb.sheetnames
        log = wb["Expense Log"]
        rows = [r for r in log.iter_rows(min_row=5, max_col=11, values_only=True) if r[0] is not None]
        eq(len(rows), 56)
        eq(round(sum(r[8] for r in rows), 2), 10891.09)
        assert len(wb["Dashboard"]._charts) == 2, "charts missing"
        assert len(log.data_validations.dataValidation) >= 9, "dropdowns missing"
        eq(wb["Lists"]["H4"].value, "Food & Groceries")
        eq(wb["Lists"]["A5"].value, "Raja")
    check("export downloads the v2 layout with all 56 rows, charts and dropdowns", verify_export)

    def verify_recalc():
        tmp = f"{WORK}/recalc.xlsx"
        shutil.copy(exported, tmp)
        out = subprocess.run(["python3", RECALC, tmp, "120"], capture_output=True, text=True).stdout
        res = json.loads(out[out.index("{"):])
        assert res.get("status") == "success", res
        wb = openpyxl.load_workbook(tmp, data_only=True)
        eq(round(wb["Monthly Summary"]["Z24"].value, 2), 10891.09)
        eq(wb["Expense Log"]["O5"].value, "OK")
    check("exported workbook recalculates with no errors and the right totals", verify_recalc)

    # ---------------------------------------------------------------- import: duplicates only
    c = import_file(page, exported)
    check("re-importing the export finds 56 duplicates, 0 new", lambda: eq(c, {"New": 0, "Duplicates": 56, "Problems": 0}))
    check("import button disabled when nothing new", lambda: expect(page.get_by_role("button", name=re.compile("^Import"))).to_be_disabled())
    page.screenshot(path=f"{SHOTS}/p2-03-import-dups.png", full_page=True)

    c = import_file(page, V2)
    check("importing the original v2 workbook: all 56 duplicates", lambda: eq(c, {"New": 0, "Duplicates": 56, "Problems": 0}))
    c = import_file(page, V1)
    check("importing the very first workbook layout: all 56 duplicates", lambda: eq(c, {"New": 0, "Duplicates": 56, "Problems": 0}))

    partial = f"{WORK}/partial.xlsx"
    wbp = openpyxl.load_workbook(V2)
    wsp = wbp["Expense Log"]
    for r in range(15, 61):  # keep the first 10 rows only
        for col in range(1, 12):
            wsp.cell(r, col).value = None
    wsp.cell(15, 1).value = dt_mod.datetime(2026, 10, 18); wsp.cell(15, 2).value = "Bunnings"
    for col, v in zip(range(3, 12), ["Housing & Rent", "Household Supplies", "Variable Wants", "Raja", "Shared (50/50)", "Card", 23.5, "No", ""]):
        wsp.cell(15, col).value = v
    wbp.save(partial)
    c = import_file(page, partial)
    check("partial overlap: 10 duplicates, 1 new", lambda: eq(c, {"New": 1, "Duplicates": 10, "Problems": 0}))

    # ---------------------------------------------------------------- import: mixed file
    mixed = f"{WORK}/mixed.xlsx"
    wb = openpyxl.load_workbook(exported)
    ws = wb["Expense Log"]
    import datetime as dt
    extra = [
        (dt.datetime(2026, 10, 18), "Officeworks", "Shopping", "Electronics", "Variable Wants", "Raja", "Shared (50/50)", "Card", 49.0, "No", ""),
        (dt.datetime(2026, 10, 18), "Myki Top Up", "Transport", "Public Transport (Myki)", "Fixed Needs", "Jo", "Jo only", "", 5.7, "No", "twice"),
        (dt.datetime(2026, 10, 18), "Myki Top Up", "Transport", "Public Transport (Myki)", "Fixed Needs", "Jo", "Jo only", "", 5.7, "No", "twice"),
        (dt.datetime(2026, 10, 12), "October Rent", "Housing & Rent", "Rent", "Fixed Needs", "Raja", "Shared (50/50)", "", 2595, "No", "dup"),
        (dt.datetime(2026, 10, 18), "Vet visit", "Pets", "Vet", "Fixed Needs", "Jo", "Shared (50/50)", "", 80, "No", ""),
    ]
    for i, row in enumerate(extra):
        for j, v in enumerate(row):
            ws.cell(61 + i, j + 1).value = v
    wb.save(mixed)
    c = import_file(page, mixed)
    check("mixed file: 4 new (incl. repeats beyond what is stored), 56 duplicates, 1 problem",
          lambda: eq(c, {"New": 4, "Duplicates": 56, "Problems": 1}))
    check("unknown category offered for creation", lambda: expect(page.get_by_text("These categories aren't in the app yet")).to_be_visible())
    page.screenshot(path=f"{SHOTS}/p2-04-import-mixed.png", full_page=True)
    page.get_by_role("button", name="Import", exact=True).click()
    check("import result message", lambda: expect(page.get_by_role("status")).to_contain_text("Imported 5 new expenses. Skipped 56 duplicates"))
    check("database now has 61 expenses", lambda: eq(live_count(), "61"))
    check("Pets category was created", lambda: eq(sql("select count(*) from categories where name='Pets'"), "1"))

    c = import_file(page, mixed)
    check("importing the same file again adds nothing", lambda: eq(c, {"New": 0, "Duplicates": 61, "Problems": 0}))

    page.goto(BASE + "/settings/data")
    page.get_by_role("button", name="Undo").first.click()
    page.wait_for_timeout(1500)
    check("undo import removes the 5 rows", lambda: eq(live_count(), "56"))

    # ---------------------------------------------------------------- recurring via Add → Repeat
    page.goto(BASE + "/add")
    page.fill("#amount", "18.99")
    page.fill("input[name=merchant]", "Netflix")
    page.fill("input[name=txn_date]", "2026-08-18")
    page.get_by_role("button", name="Lifestyle & Entertainment").click()
    page.get_by_role("button", name="Subscriptions").click()
    page.get_by_role("button", name="Variable Wants").click()
    page.select_option("select[name=repeat]", "monthly")
    page.get_by_role("button", name="Save", exact=True).click()
    expect(page).to_have_url(re.compile(r"/transactions\?month=2026-08"))
    check("repeat creates the expense plus missed months (Aug, Sep, Oct)", lambda: eq(
        sql("select string_agg(txn_date::text, ',' order by txn_date) from transactions where merchant='Netflix' and deleted_at is null"),
        "2026-08-18,2026-09-18,2026-10-18"))
    check("recurring badge in list", lambda: expect(page.get_by_text("RECURRING").first).to_be_visible())
    for _ in range(3):
        page.goto(BASE + "/")
    check("opening the app again doesn't duplicate", lambda: eq(
        sql("select count(*) from transactions where merchant='Netflix'"), "3"))

    page.goto(BASE + "/settings/recurring")
    check("recurring list shows next date 18 Nov", lambda: expect(page.locator("li").filter(has_text="Netflix")).to_contain_text("Next: Wed, 18 Nov"))

    # pause / resume
    page.locator("li").filter(has_text="Netflix").get_by_role("button", name="Pause").click()
    check("pause works", lambda: expect(page.locator("li").filter(has_text="Netflix")).to_contain_text("Paused"))
    page.locator("li").filter(has_text="Netflix").get_by_role("button", name="Resume").click()
    check("resume works", lambda: expect(page.locator("li").filter(has_text="Netflix")).to_contain_text("Next: Wed, 18 Nov"))

    # edit amount: future only
    page.locator("li").filter(has_text="Netflix").get_by_role("link").click()
    page.fill("#amount", "22.99")
    page.get_by_role("button", name="Save recurring expense").click()
    expect(page).to_have_url(re.compile("/settings/recurring$"))
    check("editing the rule changes future only", lambda: (
        expect(page.locator("li").filter(has_text="Netflix")).to_contain_text("$22.99"),
        eq(sql("select string_agg(distinct amount::text, ',') from transactions where merchant='Netflix'"), "18.99")))

    # ---------------------------------------------------------------- recurring from an existing expense
    page.goto(BASE + "/transactions?month=2026-10&q=October+Rent")
    page.get_by_role("link", name=re.compile("October Rent")).click()
    page.get_by_role("link", name="Make this a recurring expense").click()
    check("form prefilled from the expense, starting next month", lambda: (
        expect(page.locator("input[name=merchant]")).to_have_value("October Rent"),
        expect(page.locator("input[name=txn_date]")).to_have_value("2026-11-12")))
    page.fill("input[name=merchant]", "Rent")
    page.get_by_role("button", name="Create recurring expense").click()
    expect(page).to_have_url(re.compile("/settings/recurring$"))
    check("future rule creates nothing yet", lambda: eq(sql("select count(*) from transactions where merchant='Rent'"), "0"))
    check("new rule listed with next date", lambda: expect(page.locator("li").filter(has_text="Rent").filter(has_text="Monthly")).to_contain_text("Next: Thu, 12 Nov"))
    page.screenshot(path=f"{SHOTS}/p2-05-recurring.png", full_page=True)

    # delete rule keeps history
    page.locator("li").filter(has_text="Netflix").get_by_role("link").click()
    page.get_by_role("button", name="Delete recurring expense").click()
    expect(page).to_have_url(re.compile("/settings/recurring$"))
    check("deleting a rule keeps the expenses it created", lambda: eq(sql("select count(*) from transactions where merchant='Netflix' and deleted_at is null"), "3"))

    page.goto(BASE + "/settings")
    page.screenshot(path=f"{SHOTS}/p2-06-settings.png", full_page=True)
    page.goto(BASE + "/")
    page.screenshot(path=f"{SHOTS}/p2-07-home.png", full_page=True)
    browser.close()

for r in results:
    print(*r)
fails = [r for r in results if r[0] == "FAIL"]
print(f"\n{len(results) - len(fails)}/{len(results)} passed")
sys.exit(1 if fails else 0)
